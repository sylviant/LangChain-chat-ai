from fpdf import FPDF
from docx import Document
import io
import streamlit as st
from langchain_groq import ChatGroq
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder
from langchain_core.messages import HumanMessage, AIMessage
from langchain_community.chat_message_histories import ChatMessageHistory
from langchain_core.output_parsers import StrOutputParser
from langchain_core.runnables import RunnablePassthrough
from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_community.tools import DuckDuckGoSearchRun
from groq import Groq
from gtts import gTTS
from dotenv import load_dotenv
from database import (
    nouvelle_conversation, ajouter_message, get_messages,
    get_conversations, supprimer_conversation, renommer_conversation,
    mettre_a_jour_titre
)
import math
import datetime
import tempfile
import os
import subprocess
import sounddevice as sd
import soundfile as sf

load_dotenv()

groq_client = Groq()
search = DuckDuckGoSearchRun()

# ---- THEME ----
def appliquer_theme(theme):
    couleur = theme["couleur_principale"]
    taille = theme["taille_police"]
    st.markdown(f"""
    <style>
        .stMarkdown, .stChatMessage, p, li {{
            font-size: {taille}px !important;
        }}
        .stButton > button {{
            border-color: {couleur} !important;
            color: {couleur} !important;
        }}
        .stButton > button:hover {{
            background-color: {couleur} !important;
            color: white !important;
        }}
        .stChatInput textarea:focus {{
            border-color: {couleur} !important;
            box-shadow: 0 0 0 1px {couleur} !important;
        }}
        .stRadio label, .stSelectbox label {{
            color: {couleur} !important;
            font-weight: 500;
        }}
        [data-testid="stSidebar"] {{
            border-right: 2px solid {couleur} !important;
        }}
        h1 {{
            color: {couleur} !important;
        }}
        [data-testid="stChatMessage"] {{
            border-left: 3px solid {couleur};
            padding-left: 8px;
        }}
    </style>
    """, unsafe_allow_html=True)

# ---- TITRE AUTO ----
def generer_titre(user_input):
    llm = ChatGroq(model="llama-3.3-70b-versatile", temperature=0)
    prompt = ChatPromptTemplate.from_messages([
        ("system", "Tu es un assistant qui génère des titres très courts (3-5 mots maximum) pour résumer une question. Réponds UNIQUEMENT avec le titre, sans ponctuation ni guillemets."),
        ("human", "{input}"),
    ])
    chain = prompt | llm | StrOutputParser()
    titre = chain.invoke({"input": user_input})
    return titre.strip()[:50]

# ---- EXPORT ----
def exporter_txt(messages, titre):
    contenu = f"Conversation : {titre}\n"
    contenu += "=" * 40 + "\n\n"
    for msg in messages:
        role = "Toi" if msg["role"] == "user" else "AI"
        contenu += f"{role} :\n{msg['content']}\n\n"
    return contenu.encode("utf-8")

def exporter_pdf(messages, titre):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, titre[:60], ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 8, f"Exporté le {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}", ln=True)
    pdf.ln(5)
    for msg in messages:
        role = "Toi" if msg["role"] == "user" else "AI"
        pdf.set_font("Helvetica", "B", 11)
        pdf.set_fill_color(230, 240, 255) if msg["role"] == "user" else pdf.set_fill_color(240, 255, 240)
        pdf.cell(0, 8, role, ln=True, fill=True)
        pdf.set_font("Helvetica", "", 10)
        texte = msg["content"].encode("latin-1", errors="replace").decode("latin-1")
        pdf.multi_cell(0, 6, texte)
        pdf.ln(3)
    return bytes(pdf.output())

def exporter_docx(messages, titre):
    doc = Document()
    doc.add_heading(titre, level=1)
    doc.add_paragraph(f"Exporté le {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}")
    for msg in messages:
        role = "Toi" if msg["role"] == "user" else "AI"
        p = doc.add_paragraph()
        p.add_run(f"{role} : ").bold = True
        p.add_run(msg["content"])
        doc.add_paragraph()
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

# ---- CONFIG ----
st.set_page_config(page_title="Chat AI", page_icon="🤖", layout="wide")

# ---- INITIALISATION SESSION ----
if "conversation_id" not in st.session_state:
    st.session_state.conversation_id = nouvelle_conversation()
if "messages" not in st.session_state:
    st.session_state.messages = []
if "rag_chain" not in st.session_state:
    st.session_state.rag_chain = None
if "pdf_loaded" not in st.session_state:
    st.session_state.pdf_loaded = False
if "mode" not in st.session_state:
    st.session_state.mode = "Chat normal"
if "personnalite" not in st.session_state:
    st.session_state.personnalite = "Assistant général"
if "theme" not in st.session_state:
    st.session_state.theme = {
        "couleur_principale": "#1f77b4",
        "taille_police": 16,
    }

# ---- PERSONNALITES ----
PERSONNALITES = {
    "Assistant général": "Tu es un assistant IA utile et bienveillant. Réponds toujours en français.",
    "Professeur": "Tu es un professeur patient et pédagogue. Tu expliques les concepts de manière claire et simple avec des exemples concrets. Réponds toujours en français.",
    "Assistant juridique": "Tu es un assistant juridique compétent. Tu expliques les concepts juridiques clairement tout en précisant que tes réponses ne remplacent pas un avocat. Réponds toujours en français.",
    "Coach sportif": "Tu es un coach sportif motivant et bienveillant. Tu donnes des conseils sur l'entraînement, la nutrition et la récupération. Réponds toujours en français.",
    "Développeur": "Tu es un développeur senior expérimenté. Tu aides avec le code, le debugging et l'architecture logicielle. Réponds toujours en français.",
    "Assistant médical": "Tu es un assistant médical informatif. Tu fournis des informations de santé claires tout en précisant que tes réponses ne remplacent pas un médecin. Réponds toujours en français.",
    "Créatif": "Tu es un assistant créatif et imaginatif. Tu aides avec l'écriture, le brainstorming et les projets artistiques. Réponds toujours en français.",
}

appliquer_theme(st.session_state.theme)

# ---- FONCTIONS CHAT ----
def get_chat_chain():
    llm = ChatGroq(model="llama-3.3-70b-versatile", temperature=0.7)
    system_prompt = PERSONNALITES.get(
        st.session_state.personnalite,
        PERSONNALITES["Assistant général"]
    )
    prompt = ChatPromptTemplate.from_messages([
        ("system", system_prompt),
        MessagesPlaceholder(variable_name="history"),
        ("human", "{input}"),
    ])
    history_messages = [
        HumanMessage(content=m["content"]) if m["role"] == "user"
        else AIMessage(content=m["content"])
        for m in st.session_state.messages
    ]
    chain = (
        RunnablePassthrough.assign(history=lambda _: history_messages)
        | prompt
        | llm
        | StrOutputParser()
    )
    return chain

# ---- FONCTIONS RAG ----
def create_rag_chain(pdf_path):
    loader = PyPDFLoader(pdf_path)
    pages = loader.load()
    splitter = RecursiveCharacterTextSplitter(chunk_size=500, chunk_overlap=50)
    docs = splitter.split_documents(pages)
    embeddings = HuggingFaceEmbeddings(model_name="sentence-transformers/all-MiniLM-L6-v2")
    vectorstore = FAISS.from_documents(docs, embeddings)
    retriever = vectorstore.as_retriever(search_kwargs={"k": 3})
    llm = ChatGroq(model="llama-3.3-70b-versatile", temperature=0)
    prompt = ChatPromptTemplate.from_messages([
        ("system", """Tu es un assistant qui répond aux questions basées uniquement sur le document fourni.
Contexte : {context}
Réponds toujours en français. Si la réponse n'est pas dans le document, dis-le clairement."""),
        ("human", "{question}"),
    ])
    chain = (
        {"context": retriever, "question": RunnablePassthrough()}
        | prompt
        | llm
        | StrOutputParser()
    )
    return chain

# ---- FONCTIONS AGENT ----
def executer_outil(nom, argument):
    if nom == "recherche_internet":
        return search.run(argument)
    elif nom == "calculatrice":
        try:
            result = eval(argument, {"__builtins__": {}}, {
                "sqrt": math.sqrt, "pi": math.pi,
                "sin": math.sin, "cos": math.cos,
                "tan": math.tan, "log": math.log,
                "abs": abs, "round": round, "pow": pow,
            })
            return str(result)
        except Exception as e:
            return f"Erreur : {e}"
    elif nom == "date_heure":
        now = datetime.datetime.now()
        return f"Nous sommes le {now.strftime('%A %d %B %Y')} et il est {now.strftime('%H:%M:%S')}."
    elif nom == "sauvegarde_note":
        with open("notes.txt", "a", encoding="utf-8") as f:
            timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
            f.write(f"[{timestamp}] {argument}\n")
        return "Note sauvegardée !"
    return "Outil inconnu."

def run_agent(user_input):
    llm = ChatGroq(model="llama-3.3-70b-versatile", temperature=0)
    prompt = ChatPromptTemplate.from_messages([
        ("system", """Tu es un assistant IA autonome et intelligent.
Tu as accès aux outils suivants. Quand tu as besoin d'un outil, réponds UNIQUEMENT avec ce format exact :

OUTIL: nom_outil
ARGUMENT: l'argument

Les outils disponibles sont :
- recherche_internet : pour chercher sur le web
- calculatrice : pour faire des calculs
- date_heure : pour la date et heure actuelle
- sauvegarde_note : pour sauvegarder une note

Quand tu as la réponse finale, réponds normalement en français sans le format OUTIL/ARGUMENT."""),
        ("human", "{input}"),
    ])
    chain = prompt | llm | StrOutputParser()
    context = user_input
    log = []
    for _ in range(4):
        response = chain.invoke({"input": context})
        if "OUTIL:" in response and "ARGUMENT:" in response:
            lines = response.strip().split("\n")
            nom_outil = ""
            argument = ""
            for line in lines:
                if line.startswith("OUTIL:"):
                    nom_outil = line.replace("OUTIL:", "").strip()
                elif line.startswith("ARGUMENT:"):
                    argument = line.replace("ARGUMENT:", "").strip()
            tool_result = executer_outil(nom_outil, argument)
            log.append(f"**Outil :** `{nom_outil}` | **Argument :** `{argument}`")
            log.append(f"**Résultat :** {tool_result}")
            context = f"""Question originale : {user_input}
Tu as utilisé l'outil '{nom_outil}' et voici le résultat :
{tool_result}
Maintenant réponds à la question en français en utilisant ce résultat."""
        else:
            return response, log
    return response, log

# ---- FONCTIONS VOCALES ----
def speak(text):
    tts = gTTS(text=text, lang="fr")
    with tempfile.NamedTemporaryFile(delete=False, suffix=".mp3") as f:
        tmp_path = f.name
    tts.save(tmp_path)
    subprocess.run(["start", "/wait", tmp_path], shell=True)
    os.unlink(tmp_path)

def listen():
    sample_rate = 16000
    duration = 5
    recording = sd.rec(
        int(duration * sample_rate),
        samplerate=sample_rate,
        channels=1,
        dtype="int16"
    )
    sd.wait()
    with tempfile.NamedTemporaryFile(delete=False, suffix=".wav") as f:
        tmp_path = f.name
    sf.write(tmp_path, recording, sample_rate)
    with open(tmp_path, "rb") as audio_file:
        transcription = groq_client.audio.transcriptions.create(
            model="whisper-large-v3",
            file=audio_file,
            language="fr"
        )
    os.unlink(tmp_path)
    return transcription.text.strip()

# ---- SIDEBAR ----
with st.sidebar:
    st.title("Chat AI")

    mode = st.radio("Mode", [
        "Chat normal",
        "Chat avec PDF",
        "Agent autonome",
        "Chat vocal"
    ])
    st.session_state.mode = mode

    st.divider()

    st.subheader("Personnalité")
    personnalite = st.selectbox(
        "Choisir un assistant",
        list(PERSONNALITES.keys()),
        index=list(PERSONNALITES.keys()).index(st.session_state.personnalite)
    )
    if personnalite != st.session_state.personnalite:
        st.session_state.personnalite = personnalite
        st.rerun()

    descriptions = {
        "Assistant général": "Assistant polyvalent",
        "Professeur": "Explique avec pédagogie",
        "Assistant juridique": "Questions de droit",
        "Coach sportif": "Sport et nutrition",
        "Développeur": "Code et debugging",
        "Assistant médical": "Informations de santé",
        "Créatif": "Écriture et créativité",
    }
    st.caption(descriptions.get(personnalite, ""))

    st.divider()

    st.subheader("Apparence")
    couleur = st.color_picker(
        "Couleur principale",
        st.session_state.theme["couleur_principale"]
    )
    taille_police = st.slider(
        "Taille de police",
        min_value=12,
        max_value=24,
        value=st.session_state.theme["taille_police"]
    )
    if couleur != st.session_state.theme["couleur_principale"] or \
            taille_police != st.session_state.theme["taille_police"]:
        st.session_state.theme["couleur_principale"] = couleur
        st.session_state.theme["taille_police"] = taille_police
        st.rerun()

    st.divider()

    if mode == "Chat avec PDF":
        uploaded_file = st.file_uploader("Uploade ton PDF", type="pdf")
        if uploaded_file and not st.session_state.pdf_loaded:
            with st.spinner("Chargement du PDF..."):
                with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                    tmp.write(uploaded_file.read())
                    tmp_path = tmp.name
                st.session_state.rag_chain = create_rag_chain(tmp_path)
                st.session_state.pdf_loaded = True
            st.success("PDF chargé !")
        st.divider()

    st.subheader("Exporter la conversation")
    if st.session_state.messages:
        conversations_list = get_conversations()
        titre_conv = next(
            (c["titre"] for c in conversations_list if c["id"] == st.session_state.conversation_id),
            "conversation"
        )
        col1, col2, col3 = st.columns(3)
        with col1:
            st.download_button(
                label="TXT",
                data=exporter_txt(st.session_state.messages, titre_conv),
                file_name=f"{titre_conv[:20]}.txt",
                mime="text/plain",
                use_container_width=True
            )
        with col2:
            st.download_button(
                label="PDF",
                data=exporter_pdf(st.session_state.messages, titre_conv),
                file_name=f"{titre_conv[:20]}.pdf",
                mime="application/pdf",
                use_container_width=True
            )
        with col3:
            st.download_button(
                label="DOCX",
                data=exporter_docx(st.session_state.messages, titre_conv),
                file_name=f"{titre_conv[:20]}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )
    else:
        st.caption("Aucun message à exporter.")

    st.divider()

    st.subheader("Conversations")
    if st.button("Nouvelle conversation", use_container_width=True):
        st.session_state.conversation_id = nouvelle_conversation()
        st.session_state.messages = []
        st.session_state.pdf_loaded = False
        st.session_state.rag_chain = None
        st.rerun()

    conversations = get_conversations()
    for conv in conversations:
        col1, col2 = st.columns([3, 1])
        with col1:
            if st.button(f"{conv['titre'][:25]} — {conv['date']}", key=f"conv_{conv['id']}", use_container_width=True):
                st.session_state.conversation_id = conv["id"]
                st.session_state.messages = get_messages(conv["id"])
                st.rerun()
        with col2:
            if st.button("🗑", key=f"del_{conv['id']}"):
                supprimer_conversation(conv["id"])
                if conv["id"] == st.session_state.conversation_id:
                    st.session_state.conversation_id = nouvelle_conversation()
                    st.session_state.messages = []
                st.rerun()

# ---- TITRE PRINCIPAL ----
titres = {
    "Chat normal": "Chat AI",
    "Chat avec PDF": "Chat avec ton PDF",
    "Agent autonome": "Agent AI autonome",
    "Chat vocal": "Chat vocal avec Whisper"
}
icones = {
    "Assistant général": "🤖",
    "Professeur": "👨‍🏫",
    "Assistant juridique": "⚖️",
    "Coach sportif": "💪",
    "Développeur": "💻",
    "Assistant médical": "🩺",
    "Créatif": "🎨",
}
titre_mode = titres.get(mode, "Chat AI")
icone = icones.get(st.session_state.personnalite, "🤖")
st.title(f"{icone} {titre_mode} — {st.session_state.personnalite}")

# ---- AFFICHAGE MESSAGES ----
for msg in st.session_state.messages:
    with st.chat_message(msg["role"]):
        st.markdown(msg["content"])

# ---- CHAT VOCAL ----
if mode == "Chat vocal":
    if st.button("Parler au micro (5 sec)", use_container_width=True):
        with st.spinner("Écoute en cours..."):
            user_input = listen()
        if user_input:
            st.session_state.messages.append({"role": "user", "content": user_input})
            ajouter_message(st.session_state.conversation_id, "user", user_input)
            if len(st.session_state.messages) == 1:
                titre = generer_titre(user_input)
                mettre_a_jour_titre(st.session_state.conversation_id, titre)
            chain = get_chat_chain()
            with st.chat_message("user"):
                st.markdown(user_input)
            with st.chat_message("assistant"):
                response = st.write_stream(chain.stream({"input": user_input}))
            ajouter_message(st.session_state.conversation_id, "assistant", response)
            st.session_state.messages.append({"role": "assistant", "content": response})
            with st.spinner("Lecture vocale..."):
                speak(response)
            st.rerun()

# ---- SAISIE TEXTE ----
if user_input := st.chat_input("Écris ton message..."):
    st.session_state.messages.append({"role": "user", "content": user_input})
    ajouter_message(st.session_state.conversation_id, "user", user_input)

    if len(st.session_state.messages) == 1:
        titre = generer_titre(user_input)
        mettre_a_jour_titre(st.session_state.conversation_id, titre)

    with st.chat_message("user"):
        st.markdown(user_input)

    with st.chat_message("assistant"):
        if mode == "Chat normal":
            chain = get_chat_chain()
            response = st.write_stream(chain.stream({"input": user_input}))

        elif mode == "Chat avec PDF":
            if st.session_state.rag_chain is None:
                st.warning("Uploade d'abord un PDF dans le menu à gauche.")
                response = None
            else:
                response = st.write_stream(
                    st.session_state.rag_chain.stream(user_input)
                )

        elif mode == "Agent autonome":
            with st.spinner("L'agent réfléchit..."):
                response, log = run_agent(user_input)
            if log:
                with st.expander("Outils utilisés"):
                    for entry in log:
                        st.markdown(entry)
            st.markdown(response)

        elif mode == "Chat vocal":
            chain = get_chat_chain()
            response = st.write_stream(chain.stream({"input": user_input}))
            with st.spinner("Lecture vocale..."):
                speak(response)

    if response:
        ajouter_message(st.session_state.conversation_id, "assistant", response)
        st.session_state.messages.append({"role": "assistant", "content": response})